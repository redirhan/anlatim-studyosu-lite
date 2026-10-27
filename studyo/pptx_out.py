"""Şekil listesi → düzenlenebilir PPTX (python-pptx). Her metin gerçek metin kutusu, diyagramlar gerçek şekil."""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .layout import SH, SW, Box, Line, Pic, Slide, Table, TextBox

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL, "rarrow": MSO_SHAPE.RIGHT_ARROW}


def rgb(c):
    return RGBColor(*[int(v) for v in c])


def _hex(c):
    return "%02X%02X%02X" % tuple(int(v) for v in c)


def _bullet_xml(p, tb: TextBox, para, size):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if not para.bullet:
        etree.SubElement(pPr, qn("a:buNone"))
        return
    ind = int(Pt(size * 1.25))
    pPr.set("marL", str(ind))
    pPr.set("indent", str(-ind))
    if tb.bullet_color:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", _hex(tb.bullet_color))
    if para.bullet == "num":
        f = etree.SubElement(pPr, qn("a:buFont"))
        f.set("typeface", "+mj-lt")
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", tb.bullet_char)


def fill_text(tf, tb: TextBox):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = ANCHOR[tb.valign]
    size = tb.size or tb.max_size
    for i, para in enumerate(tb.paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        s = size * para.rel
        p.alignment = ALIGN[para.align]
        p.line_spacing = tb.spacing
        p.space_after = Pt(s * para.space_after) if i + 1 < len(tb.paras) else Pt(0)
        _bullet_xml(p, tb, para, s)
        for r in para.runs:
            run = p.add_run()
            run.text = r.text
            f = run.font
            f.name = tb.face
            f.size = Pt(s)
            f.bold = r.bold
            f.italic = r.italic
            f.color.rgb = rgb(r.color or tb.color)


def _no_shadow(shape):
    """Tema efektlerini (gölge) kaldır: boş effectLst + p:style'ı sil (LibreOffice effectRef'e de bakıyor)."""
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _cell_border(cell, color, w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(Pt(w_pt))))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", _hex(color))


def add_shapes(slide, shapes):
    for sh in shapes:
        if isinstance(sh, TextBox):
            box = slide.shapes.add_textbox(Inches(sh.x), Inches(sh.y), Inches(sh.w), Inches(sh.h))
            box.name = sh.name
            fill_text(box.text_frame, sh)
        elif isinstance(sh, Box):
            kind = SHAPES[sh.shape]
            if sh.shape == "rect" and sh.radius > 0:
                kind = MSO_SHAPE.ROUNDED_RECTANGLE
            s = slide.shapes.add_shape(kind, Inches(sh.x), Inches(sh.y), Inches(sh.w), Inches(sh.h))
            s.name = sh.name
            if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                s.adjustments[0] = min(0.5, sh.radius / max(0.01, min(sh.w, sh.h)))
            if sh.fill is not None:
                s.fill.solid()
                s.fill.fore_color.rgb = rgb(sh.fill)
            else:
                s.fill.background()
            if sh.line is not None:
                s.line.color.rgb = rgb(sh.line)
                s.line.width = Pt(sh.line_w)
            else:
                s.line.fill.background()
            _no_shadow(s)
            if sh.text is not None:
                fill_text(s.text_frame, sh.text)
                s.text_frame.margin_left = s.text_frame.margin_right = Inches(0)
        elif isinstance(sh, Line):
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(sh.x1), Inches(sh.y1), Inches(sh.x2),
                                           Inches(sh.y2))
            c.name = sh.name
            c.line.color.rgb = rgb(sh.color)
            c.line.width = Pt(sh.width)
        elif isinstance(sh, Pic):
            p = slide.shapes.add_picture(sh.path, Inches(sh.x), Inches(sh.y), Inches(sh.w), Inches(sh.h))
            p.name = sh.name
            if sh.alt:
                p._element.nvPicPr.cNvPr.set("descr", sh.alt)
        elif isinstance(sh, Table):
            nr, nc = len(sh.rows), len(sh.col_w)
            gf = slide.shapes.add_table(nr, nc, Inches(sh.x), Inches(sh.y), Inches(sh.w), Inches(sh.h))
            gf.name = sh.name
            tbl = gf.table
            tbl.first_row = True
            tbl.horz_banding = False
            for j, w in enumerate(sh.col_w):
                tbl.columns[j].width = Inches(w)
            for i, h in enumerate(sh.row_h):
                tbl.rows[i].height = Inches(h)
            for i, row in enumerate(sh.rows):
                for j in range(nc):
                    cell = tbl.cell(i, j)
                    txt = str(row[j]) if j < len(row) else ""
                    head = i == 0
                    lab = j == 0 and not head
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(sh.head_fill if head else sh.label_fill if lab else (255, 255, 255)
                                                   if sum(sh.ink) < 384 else (28, 37, 56))
                    tb = TextBox(0, 0, 1, 1, [], sh.face, sh.head_ink if head else sh.ink, sh.size, sh.size)
                    from .layout import Para, Run
                    tb.paras = [Para([Run(txt, head or lab)])]
                    tb.size = sh.size
                    tb.valign = "m"
                    fill_text(cell.text_frame, tb)
                    cell.margin_left = cell.margin_right = Inches(0.1)
                    cell.margin_top = cell.margin_bottom = Inches(0.05)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    _cell_border(cell, sh.line)


def write(slides: list[Slide], out: Path, meta: dict):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]
    for sl in slides:
        s = prs.slides.add_slide(blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(sl.bg)
        add_shapes(s, sl.shapes)
        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes
    cp = prs.core_properties
    cp.title = meta.get("title", "")
    cp.author = meta.get("author", "") or ""
    cp.subject = meta.get("subtitle", "") or ""
    cp.keywords = ", ".join(meta.get("keywords", []))
    cp.comments = "Anlatım Stüdyosu ile hazırlandı; tüm metinler düzenlenebilir."
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
