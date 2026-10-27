"""Kullanıcının kendi notları/kaynakları → düz metin (araştırma ve doğrulama için).

Proje klasörü:
  notes/          PDF, DOCX, PPTX, TXT/MD, görsel (OCR yalnız tesseract kuruluysa; yoksa görseli Claude okur)
Çıktı:
  work/notes.txt        tüm metin, dosya başlıklarıyla ("=== dosya ===")
  work/notes_index.json {dosya: {"chars": n, "method": "...", "warning": "..."}}
"""
from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

TEXT_EXT = {".txt", ".md"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".heic", ".webp", ".tif", ".tiff"}


def _pdf(p):
    from pypdf import PdfReader
    r = PdfReader(str(p))
    pages = [(pg.extract_text() or "") for pg in r.pages]
    return "\n\n".join(f"[s. {i + 1}]\n{t}" for i, t in enumerate(pages)), f"pypdf, {len(pages)} sayfa"


def _docx(p):
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for tb in d.tables:
        for row in tb.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts), "python-docx"


def _pptx(p):
    from pptx import Presentation
    pr = Presentation(str(p))
    out = []
    for i, sl in enumerate(pr.slides, 1):
        txt = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                txt += [para.text for para in sh.text_frame.paragraphs if para.text.strip()]
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    txt.append(" | ".join(c.text for c in row.cells))
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame is not None:
            n = sl.notes_slide.notes_text_frame.text.strip()
            if n:
                txt.append("Not: " + n)
        out.append(f"[slayt {i}]\n" + "\n".join(txt))
    return "\n\n".join(out), f"python-pptx, {len(pr.slides)} slayt"


def _image(p):
    if shutil.which("tesseract"):
        r = subprocess.run(["tesseract", str(p), "-", "-l", "tur+eng"], capture_output=True, text=True)
        return r.stdout, "tesseract OCR"
    return "", "OCR yok (tesseract kurulu değil) — elle yaz veya fotoğraftan metni Claude okusun"


def _other(p):
    if shutil.which("textutil"):   # macOS: .doc, .rtf, .odt
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(p)], capture_output=True, text=True)
        if r.returncode == 0:
            return r.stdout, "textutil"
    return "", "desteklenmeyen biçim"


def extract(proj: Path) -> dict:
    notes_dir = proj / "notes"
    work = proj / "work"
    work.mkdir(exist_ok=True)
    index, chunks = {}, []
    for p in sorted(notes_dir.rglob("*")):
        if not p.is_file() or p.name.startswith("."):
            continue
        ext = p.suffix.lower()
        try:
            if ext in TEXT_EXT:
                txt, how = p.read_text(encoding="utf-8", errors="replace"), "metin"
            elif ext == ".pdf":
                txt, how = _pdf(p)
            elif ext == ".docx":
                txt, how = _docx(p)
            elif ext == ".pptx":
                txt, how = _pptx(p)
            elif ext in IMG_EXT:
                txt, how = _image(p)
            else:
                txt, how = _other(p)
        except Exception as e:  # bozuk dosya işi durdurmasın, rapora düşsün
            txt, how = "", f"hata: {e}"
        rel = str(p.relative_to(notes_dir))
        warn = "metin çok az — taranmış PDF/görsel olabilir" if len(txt.strip()) < 200 else ""
        index[rel] = {"chars": len(txt), "method": how, "warning": warn}
        chunks.append(f"=== {rel} ===\n{txt.strip()}\n")
    (work / "notes.txt").write_text("\n".join(chunks), encoding="utf-8")
    (work / "notes_index.json").write_text(json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")
    return index
